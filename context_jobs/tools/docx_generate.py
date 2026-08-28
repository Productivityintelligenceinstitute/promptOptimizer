"""Generate DOCX artifacts from a structured document spec."""

from __future__ import annotations

from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from context_jobs.tools.artifact_paths import resolve_artifact_path
from context_jobs.tools.base import BaseTool

NAVY = RGBColor(0x1A, 0x3C, 0x6E)


class DocxGenerateTool(BaseTool):
    id = "docx-generate"
    name = "DOCX Generate"
    description = (
        "Generate a Word (.docx) file from a structured document spec "
        "(title, sections, optional tables). Output is always .docx."
    )
    requires_external_key = False
    input_schema = {
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Relative output path ending in .docx, e.g. deliverables/report.docx",
            },
            "title": {"type": "string"},
            "subtitle": {"type": "string"},
            "sections": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "body": {"type": "string"},
                    },
                    "required": ["heading", "body"],
                },
            },
            "tables": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "headers": {"type": "array", "items": {"type": "string"}},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array", "items": {"type": "string"}},
                        },
                    },
                },
            },
        },
        "required": ["path", "title", "sections"],
    }

    async def execute(
        self,
        arguments: dict[str, Any],
        api_key: Optional[str] = None,
        *,
        owner: str | None = None,
        run_id: str | None = None,
    ) -> str:
        if not owner or not run_id:
            raise ValueError("docx-generate requires run context (owner and run_id)")
        path_arg = str(arguments.get("path", "")).strip()
        if not path_arg.lower().endswith(".docx"):
            raise ValueError("path must end with .docx")

        target = resolve_artifact_path(owner, run_id, path_arg)
        target.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        title = str(arguments.get("title", "")).strip()
        subtitle = str(arguments.get("subtitle") or "").strip()

        if title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = NAVY
        if subtitle:
            sp = doc.add_paragraph()
            sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sr = sp.add_run(subtitle)
            sr.italic = True
            sr.font.size = Pt(11)

        for section in arguments.get("sections") or []:
            if not isinstance(section, dict):
                continue
            heading = str(section.get("heading", "")).strip()
            body = str(section.get("body", "")).strip()
            if heading:
                doc.add_heading(heading, level=2)
            if body:
                doc.add_paragraph(body)

        for table_spec in arguments.get("tables") or []:
            if not isinstance(table_spec, dict):
                continue
            headers = [str(h) for h in (table_spec.get("headers") or [])]
            rows = table_spec.get("rows") or []
            if not headers:
                continue
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Table Grid"
            for i, header in enumerate(headers):
                table.rows[0].cells[i].text = header
            for ri, row in enumerate(rows):
                if not isinstance(row, list):
                    continue
                for ci, val in enumerate(row[: len(headers)]):
                    table.rows[ri + 1].cells[ci].text = str(val)

        doc.save(str(target))
        from context_jobs.artifacts.blob_store import persist_run_artifact_file

        persist_run_artifact_file(
            owner,
            run_id,
            path_arg,
            target,
            tool_id=self.id,
        )
        return f"Generated DOCX: {path_arg}. artifact_path={target.as_posix()}"
