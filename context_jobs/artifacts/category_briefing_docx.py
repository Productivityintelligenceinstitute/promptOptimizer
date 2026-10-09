"""Turn a finished Category Strategy run into a Word briefing stakeholders can share."""

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BRIEFING_FILENAME = "category-strategy-briefing.docx"
BRIEFING_PURPOSE = "category_briefing"
NAVY = RGBColor(0x1A, 0x3C, 0x6E)

_HEADING = re.compile(r"^\s{0,3}#{1,3}\s+(.+?)\s*$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def is_category_strategy_job(job: Any) -> bool:
    name = str(getattr(job, "name", "") or "")
    template = str(getattr(job, "output_template", "") or "")
    if name.startswith("Category Strategy Intelligence Briefing"):
        return True
    return "Strategic Pulse" in template and "What Materially Changed" in template


def build_category_briefing_docx(output_text: str, *, job_name: str = "") -> bytes:
    """Build a shareable briefing from the run's primary result. Does not call a model."""
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Category Strategy Intelligence Briefing")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = NAVY

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Shareable briefing for category stakeholders")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)

    if job_name.strip():
        job_line = document.add_paragraph()
        job_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        job_line.add_run(job_name.strip()).font.size = Pt(11)

    sections = _sections_from_output(output_text or "")
    if not sections:
        sections = [("Briefing", "Not available in the run result.")]

    for heading, body in sections:
        document.add_heading(heading, level=2)
        for block in body.split("\n"):
            line = block.strip()
            if not line:
                continue
            paragraph = document.add_paragraph()
            if line.startswith(("- ", "* ")):
                paragraph.add_run("• " + _plain(line[2:]))
            else:
                paragraph.add_run(_plain(line))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sections_from_output(text: str) -> list[tuple[str, str]]:
    stripped = text.strip()
    if stripped.startswith("{") or stripped.startswith("```"):
        parsed = _json_sections(stripped)
        if parsed:
            return parsed

    sections: list[tuple[str, list[str]]] = []
    current_heading = ""
    current_lines: list[str] = []

    def flush() -> None:
        body = "\n".join(current_lines).strip()
        if current_heading or body:
            sections.append((current_heading or "Briefing", current_lines[:]))

    for raw in text.splitlines():
        match = _HEADING.match(raw)
        if match:
            if current_heading or any(line.strip() for line in current_lines):
                flush()
            current_heading = _plain(match.group(1)).rstrip(":").strip() or "Briefing"
            current_lines = []
            continue
        current_lines.append(raw)

    if current_heading or any(line.strip() for line in current_lines):
        flush()

    cleaned: list[tuple[str, str]] = []
    for heading, lines in sections:
        body = "\n".join(lines).strip()
        if heading == "Briefing" and not body:
            continue
        if not heading and not body:
            continue
        cleaned.append((heading or "Briefing", body))
    return cleaned


def _json_sections(text: str) -> list[tuple[str, str]]:
    payload = text.strip()
    if payload.startswith("```"):
        payload = payload.strip("`")
        payload = payload.split("\n", 1)[-1]
        if payload.endswith("```"):
            payload = payload[: payload.rfind("```")]
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        return []
    if not isinstance(data, dict):
        return []

    labels = {
        "strategicPulse": "Strategic Pulse",
        "materialChanges": "What Materially Changed",
        "priorityIssues": "Priority Issues",
        "emergingInsight": "Emerging Insight",
        "valueProjects": "Value-Producing Projects",
        "decisionsRequired": "Decisions Required",
        "previousActions": "Previous Actions and Learning",
        "evidenceGaps": "Evidence Gaps",
    }
    sections: list[tuple[str, str]] = []
    for key, label in labels.items():
        if key not in data:
            continue
        rendered = _render_json_value(data[key])
        if rendered:
            sections.append((label, rendered))
    return sections


def _render_json_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        lines: list[str] = []
        for item in value:
            if isinstance(item, dict):
                parts = []
                for field, content in item.items():
                    if content in (None, "", [], {}):
                        continue
                    parts.append(f"{field}: {content}")
                if parts:
                    lines.append("- " + "; ".join(parts))
            elif item not in (None, ""):
                lines.append(f"- {item}")
        return "\n".join(lines)
    if isinstance(value, dict):
        return "\n".join(f"- {key}: {content}" for key, content in value.items() if content not in (None, ""))
    return str(value).strip()


def _plain(text: str) -> str:
    return _BOLD.sub(r"\1", text).replace("`", "").strip()
